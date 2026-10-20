import competition_info
import pandas as pd
from collections import defaultdict, Counter
import math
from openpyxl import load_workbook
from datetime import datetime, timedelta
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font

color_bank = {
    # --- LIGHT COLORS (Originals) ---
    "Pastel Red": [{"hex": "#ffadad", "text": "black"}, {"hex": "#ffc2c2", "text": "black"}, {"hex": "#ffd6d6", "text": "black"}, {"hex": "#ffebeb", "text": "black"}, {"hex": "#fff5f5", "text": "black"}],
    "Apricot": [{"hex": "#ffd6a5", "text": "black"}, {"hex": "#ffdeb3", "text": "black"}, {"hex": "#ffe6c1", "text": "black"}, {"hex": "#ffeed0", "text": "black"}, {"hex": "#fff5de", "text": "black"}],
    "Maize": [{"hex": "#fdffb6", "text": "black"}, {"hex": "#feffc9", "text": "black"}, {"hex": "#feffdb", "text": "black"}, {"hex": "#ffffed", "text": "black"}, {"hex": "#fffffa", "text": "black"}],
    "Mint": [{"hex": "#caffbf", "text": "black"}, {"hex": "#d8ffcf", "text": "black"}, {"hex": "#e6ffdf", "text": "black"}, {"hex": "#f3ffef", "text": "black"}, {"hex": "#f9fff5", "text": "black"}],
    "Tiffany Blue": [{"hex": "#9bf6ff", "text": "black"}, {"hex": "#aff8ff", "text": "black"}, {"hex": "#c4faff", "text": "black"}, {"hex": "#d9fbff", "text": "black"}, {"hex": "#edfdff", "text": "black"}],
    "Sky Blue": [{"hex": "#a0c4ff", "text": "black"}, {"hex": "#b3d1ff", "text": "black"}, {"hex": "#c6deff", "text": "black"}, {"hex": "#d9ebff", "text": "black"}, {"hex": "#ecf7ff", "text": "black"}],
    "Lavender": [{"hex": "#bdb2ff", "text": "black"}, {"hex": "#ccc2ff", "text": "black"}, {"hex": "#dbd1ff", "text": "black"}, {"hex": "#eae1ff", "text": "black"}, {"hex": "#f9f0ff", "text": "black"}],
    "Mauve": [{"hex": "#ffc6ff", "text": "black"}, {"hex": "#ffd1ff", "text": "black"}, {"hex": "#ffdbff", "text": "black"}, {"hex": "#ffe6ff", "text": "black"}, {"hex": "#fff0ff", "text": "black"}],
    "Champagne": [{"hex": "#f2e9e4", "text": "black"}, {"hex": "#f7ede2", "text": "black"}, {"hex": "#f5ebe0", "text": "black"}, {"hex": "#e3d5ca", "text": "black"}, {"hex": "#d5bdaf", "text": "black"}],
    "Sage": [{"hex": "#e9edc9", "text": "black"}, {"hex": "#f0f3da", "text": "black"}, {"hex": "#f7f9eb", "text": "black"}, {"hex": "#d9e3b1", "text": "black"}, {"hex": "#ccd5ae", "text": "black"}],
    "Honey": [{"hex": "#faedcd", "text": "black"}, {"hex": "#fefae0", "text": "black"}, {"hex": "#e9edc9", "text": "black"}, {"hex": "#f4f1de", "text": "black"}, {"hex": "#e0e1dd", "text": "black"}],
    "Peach": [{"hex": "#fec89a", "text": "black"}, {"hex": "#f9dcc4", "text": "black"}, {"hex": "#f8edeb", "text": "black"}, {"hex": "#fcd5ce", "text": "black"}, {"hex": "#ffb5a7", "text": "black"}],

    # --- UPDATED COLORS (Converted to Light Tints) ---
    "Crimson": [{"hex": "#ffccd5", "text": "black"}, {"hex": "#ffb3c1", "text": "black"}, {"hex": "#ff8fa3", "text": "black"}, {"hex": "#ff758f", "text": "black"}, {"hex": "#ff4d6d", "text": "black"}],
    "Deep Navy": [{"hex": "#e0e1dd", "text": "black"}, {"hex": "#778da9", "text": "black"}, {"hex": "#a3b18a", "text": "black"}, {"hex": "#caf0f8", "text": "black"}, {"hex": "#90e0ef", "text": "black"}],
    "Forest Green": [{"hex": "#d8f3dc", "text": "black"}, {"hex": "#b7e4c7", "text": "black"}, {"hex": "#95d5b2", "text": "black"}, {"hex": "#74c69d", "text": "black"}, {"hex": "#b2d8d8", "text": "black"}],
    "Royal Purple": [{"hex": "#e0aaff", "text": "black"}, {"hex": "#c77dff", "text": "black"}, {"hex": "#d7bce8", "text": "black"}, {"hex": "#e2cfea", "text": "black"}, {"hex": "#f1e4f3", "text": "black"}],
    "Chocolate": [{"hex": "#f5ebe0", "text": "black"}, {"hex": "#e3d5ca", "text": "black"}, {"hex": "#d5bdaf", "text": "black"}, {"hex": "#edede9", "text": "black"}, {"hex": "#d6ccc2", "text": "black"}],
    "Maroon": [{"hex": "#f7d1d1", "text": "black"}, {"hex": "#eeabb1", "text": "black"}, {"hex": "#e5989b", "text": "black"}, {"hex": "#ffb4a2", "text": "black"}, {"hex": "#ffcdb2", "text": "black"}],
    "Midnight Teal": [{"hex": "#e0f2f1", "text": "black"}, {"hex": "#b2dfdb", "text": "black"}, {"hex": "#80cbc4", "text": "black"}, {"hex": "#b2ebf2", "text": "black"}, {"hex": "#e1f5fe", "text": "black"}],
    "Plum": [{"hex": "#e7c6ff", "text": "black"}, {"hex": "#ffd6ff", "text": "black"}, {"hex": "#f7d9c4", "text": "black"}, {"hex": "#faedcb", "text": "black"}, {"hex": "#dbcdf0", "text": "black"}],
    "Charcoal": [{"hex": "#f8f9fa", "text": "black"}, {"hex": "#e9ecef", "text": "black"}, {"hex": "#dee2e6", "text": "black"}, {"hex": "#ced4da", "text": "black"}, {"hex": "#adb5bd", "text": "black"}],
    "Deep Olive": [{"hex": "#f0f3bd", "text": "black"}, {"hex": "#ecf39e", "text": "black"}, {"hex": "#d8e2dc", "text": "black"}, {"hex": "#ffe5d9", "text": "black"}, {"hex": "#eaf4f4", "text": "black"}],
    "Burnt Orange": [{"hex": "#fcd2af", "text": "black"}, {"hex": "#fcc29a", "text": "black"}, {"hex": "#ffb563", "text": "black"}, {"hex": "#ffe5b4", "text": "black"}, {"hex": "#ffdab9", "text": "black"}],
    "Indigo": [{"hex": "#e0e7ff", "text": "black"}, {"hex": "#c7d2fe", "text": "black"}, {"hex": "#a5b4fc", "text": "black"}, {"hex": "#c084fc", "text": "black"}, {"hex": "#e9d5ff", "text": "black"}],
    "Slate Blue": [{"hex": "#caf0f8", "text": "black"}, {"hex": "#ade8f4", "text": "black"}, {"hex": "#90e0ef", "text": "black"}, {"hex": "#bde0fe", "text": "black"}, {"hex": "#a2d2ff", "text": "black"}]
}

def assign_event_colors(event_list, color_bank):
    """
    Returns a dict: {event_name: (hex_code, text_color)}
    Groups by the second word of the event name.
    """
    event_colors = {}
    family_names = list(color_bank.keys())
    
    # 1. Map each unique second word to a color family
    unique_categories = []
    for event in event_list:
        parts = event.split()
        if len(parts) >= 2:
            cat = parts[1]
            if cat not in unique_categories:
                unique_categories.append(cat)

    category_to_family = {}
    for i, category in enumerate(unique_categories):
        if i >= len(family_names):
            print(f"‼️ WARNING: Out of unique colors! Category '{category}' is circling back to {family_names[i % 25]}.")
        
        category_to_family[category] = family_names[i % 25]

    # 2. Assign the specific hue (0-4) to each event
    category_usage_count = {cat: 0 for cat in unique_categories}

    for event in event_list:
        parts = event.split()
        if len(parts) < 2:
            print(f"SKIP: '{event}' does not have a second word.")
            continue
            
        category = parts[1]
        family_name = category_to_family[category]
        hue_idx = category_usage_count[category]
        
        if hue_idx >= 5:
            print(f"‼️ WARNING: Category '{category}' has > 5 events! Event '{event}' is circling back to hue {hue_idx % 5}.")
        
        # Extract the data from the bank
        color_data = color_bank[family_name][hue_idx % 5]
        
        # Store as a tuple: (hex, text_color)
        event_colors[event] = (color_data['hex'], color_data['text'])
        
        category_usage_count[category] += 1
            
    return event_colors

def parse_competition_excel(file_path):
    # Read Excel (no header, since layout is irregular)
    try:
        # Attempt to read the specific columns
        df = pd.read_excel(file_path, header=None, usecols="A:D")
    except FileNotFoundError:
        print(f"❌ ERROR: The file '{file_path}' was not found!")
        return {}, {}, {}

    
    event_players = defaultdict(set)   # event -> set of players
    player_events = defaultdict(set)   # player -> set of events

    current_event = None
    current_player = None

    for idx, row in df.iterrows():
        col_a, _, col_c, col_d = row[0], row[1], row[2], row[3]

        # Detect event in column A
        if pd.notna(col_a):
            cell_text = str(col_a).strip()
            if cell_text.lower().endswith("main draw"):     # we ignore all events that are not the main draw
                words = str(col_a)[11:].split()
                current_event = " ".join(words[:2])
                current_player = None
            else:
                current_event = None
            continue

        # Detect player name in column C
        if pd.notna(col_c):

            if str(col_c).strip().lower() == "name":
                current_player = None
                continue
            if str(col_c).strip().lower() == "<partner wanted>":
                current_player = None
                continue

            current_player = str(col_c).strip()
            if current_event:
                event_players[current_event].add(current_player)
                player_events[current_player].add(current_event)

    # Build cross-event participation summary
    event_correlations = {}
    for event, players in event_players.items():
        counter = Counter()
        for player in players:
            for other_event in player_events[player]:
                if other_event != event:
                    counter[other_event] += 1
        event_correlations[event] = list(counter.items())

    return event_correlations, event_players, player_events

def analyze_collisions(event_correlations):
    strong_collision = {}
    weak_collision = {}
    super_weak_collision = {}

    for event, correlations in event_correlations.items():
        event_parts = event.split()
        event_suffix = event_parts[1] if len(event_parts) > 1 else ""
        same = []
        different_1 = []
        different_2 = []

        for other_event, count in correlations:
            other_parts = other_event.split()
            other_suffix = other_parts[1] if len(other_parts) > 1 else ""

            if other_suffix == event_suffix:
                same.append(other_event)
            else:
                if count > 2:
                    different_1.append((other_event, count))
                else:
                    different_2.append((other_event, count))


        strong_collision[event] = same
        weak_collision[event] = different_1
        super_weak_collision[event] = different_2

    return strong_collision, weak_collision, super_weak_collision



def calculate_tournament_structure(event_players, pools):
    event_structures = {}
    round_schedules = {}

    for event, players in event_players.items():
        player_count = len(players)
        event_type = event[1].upper() if len(event) > 1 else "S"
        participants = player_count // (2 if event_type == "D" else 1)

        if participants <= 1:
            event_structures[event] = {"round_robin": [], "elimination": [], "lower_playoff": []}
            round_schedules[event] = []
            continue

        groups = []
        rr_rounds = []
        upper_elim = []
        lower_elim = []

        # --- 1. Determine Structure (Force RR for 3 or 4 participants) ---
        # If user requested pools OR if there are exactly 3 or 4 participants
        only_pool_class = participants in [3, 4]
        use_round_robin = pools or only_pool_class

        if use_round_robin:
            if participants <= 5:
                only_pool_class = True
                groups = [participants]
            else:
                # Standard partitioning for larger pools
                best_groups = None
                min_matches = float('inf')
                for y in range((participants // 4) + 1):
                    rem = participants - (y * 4)
                    if rem >= 0 and rem % 3 == 0:
                        x = rem // 3
                        total_m = (x * 3) + (y * 6)
                        if total_m < min_matches:
                            min_matches = total_m
                            best_groups = [3] * x + [4] * y
                groups = best_groups

            # Calculate RR Matches per Round
            if groups:
                max_rr_rounds = max((g if g % 2 != 0 else g - 1) for g in groups)
                for r in range(max_rr_rounds):
                    m = sum(g // 2 for g in groups if r < (g if g % 2 != 0 else g - 1))
                    rr_rounds.append(m)

        # --- 2. Elimination Logic ---
        is_elite_class = any(cls in event for cls in [" A", " V"])
        is_bronze_event = any(event.endswith(s) for s in ["9", "11", "13"])

        # If we used Round Robin and have more than 4 participants, 
        # or if pools was explicitly True, we calculate playoffs.
        # EXCEPTION: If we forced RR for 3 or 4 participants and pools was False, 
        # we stop at the RR (no elimination rounds).
        if use_round_robin:
            if only_pool_class:
                # This handles the "3 or 4 participants" case when pools=False AND "5 or less participants" when pool=True
                # They get their RR, and upper_elim stays empty.
                upper_elim = []
            else:
                # This makes elimination for classes with more than 5 participants
                advancing_per_group = 1 if is_elite_class else 2
                upper_count = len(groups) * advancing_per_group
                upper_elim = self_calculate_elim(upper_count, has_bronze=is_bronze_event)
                
                if is_bronze_event:
                    lower_count = participants - upper_count
                    if lower_count > 1:
                        lower_elim = self_calculate_elim(lower_count)
                
        else:
            # Straight elimination for everyone else
            upper_elim = self_calculate_elim(participants, has_bronze=is_bronze_event)

        # --- 3. Merge Playoff Rounds ---
        playoff_rounds = []
        num_playoff_rounds = max(len(upper_elim), len(lower_elim))
        for i in range(num_playoff_rounds):
            m_upper = upper_elim[i] if i < len(upper_elim) else 0
            m_lower = lower_elim[i] if i < len(lower_elim) else 0
            playoff_rounds.append(m_upper + m_lower)

        # --- 4. Package Results ---
        event_structures[event] = {
            "round_robin": groups,
            "elimination": upper_elim,
            "lower_playoff": lower_elim
        }
        round_schedules[event] = rr_rounds + playoff_rounds

    return event_structures, round_schedules

def self_calculate_elim(n, has_bronze=False):
    """Calculates elimination matches, adding 1 to the final round for Bronze match."""
    if n <= 1: return []
    rounds = []
    
    # Calculate the number of rounds needed (power of 2)
    x = math.ceil(math.log2(n))
    
    # First round (Handle play-ins/byes)
    # If n is a perfect power of 2, matches = n/2. Otherwise, n - next lowest power of 2.
    if 2**x == n:
        first = n // 2
    else:
        first = int(n - 2**(x - 1))
    
    rounds.append(first)
    
    # Subsequent rounds (halving matches until we hit 1)
    # The number of players remaining after the first round is 2^(x-1)
    remaining_players = 2**(x - 1)
    matches = remaining_players // 2
    
    while matches >= 1:
        rounds.append(matches)
        matches //= 2
    
    # Add Bronze Match to the final round if applicable
    if has_bronze and n >= 4 and len(rounds) > 0:
        rounds[-1] += 1
        
    return rounds

def check_rule_violations(player_events):
    warnings = []

    for player, events in player_events.items():
        events = list(events)

        # Rule 1: If playing a V event, no other event with same prefix
        '''
        v_prefixes = {e[:2] for e in events if e.endswith("V")}
        for prefix in v_prefixes:
            conflicting = [e for e in events if e.startswith(prefix) and not e.endswith("V")]
            if conflicting:
                warnings.append(
                    f"⚠️ {player}: plays {', '.join([e for e in events if e.startswith(prefix) and e.endswith('V')])}, "
                    f"so cannot play {', '.join(conflicting)}"
                )
        '''

        # Rule 2: No more than 3 total events
        if len(events) > 3:
            warnings.append(f"⚠️ {player}: participates in {len(events)} events ({', '.join(events)})")

        # Rule 3: No more than two non-V events with same prefix
        '''
        non_v_events = [e for e in events if not e.endswith("V")]
        prefix_counts = {}
        for e in non_v_events:
            prefix = e[:2]
            prefix_counts[prefix] = prefix_counts.get(prefix, 0) + 1

        for prefix, count in prefix_counts.items():
            if count > 2:
                same_prefix_events = [e for e in non_v_events if e.startswith(prefix)]
                warnings.append(
                    f"⚠️ {player}: plays {count} events with prefix '{prefix}' ({', '.join(same_prefix_events)})"
                )
        '''

        # Rule 4: No 3+ events with second character D
        '''
        d_class_events = [e for e in events if len(e) > 1 and e[1].upper() == "D"]
        if len(d_class_events) >= 3:
            warnings.append(
                f"⚠️ {player}: plays {len(d_class_events)} doubles classes ({', '.join(d_class_events)})"
            )
        '''

    return warnings

def find_problematic_players(player_events, event_players, event_correlations):
    problematic_players = {}
    checked_pairs = set()  # to avoid duplicate collision reports
    
    # Sole collision cause
    for event, collisions in event_correlations.items():
        for other_event, _ in collisions:
            # Normalize pair so ("MS V", "XD A") == ("XD A", "MS V")
            pair_key = tuple(sorted([event, other_event]))
            if pair_key in checked_pairs:
                continue  # skip duplicate
            checked_pairs.add(pair_key)

            players_in_both = event_players.get(event, set()) & event_players.get(other_event, set())

            if len(players_in_both) == 1:
                sole_player = next(iter(players_in_both))
                problematic_players.setdefault(sole_player, []).append(
                    f"Sole reason {pair_key[0]} collides with {pair_key[1]}"
                )


    return problematic_players

def write_player_events_to_excel(file_path, player_events, problematic_players, sheet_name="Player Events"):
    """
    Adds/updates a sheet listing all players, their events, and warning reasons if they're problematic.
    """
    try:
        book = load_workbook(file_path)
    except FileNotFoundError:
        print(f"❌ ERROR: The file '{file_path}' was not found. Skipping writing the player events to Excel.")
        return

    # Create sheet if missing
    if sheet_name not in book.sheetnames:
        sheet = book.create_sheet(sheet_name)
        headers = ["Player", "Events", "Warnings"]
        sheet.append(headers)

    sheet = book[sheet_name]

    # Clear only the data area (keep headers & formatting)
    for row in sheet.iter_rows(min_row=2, max_col= 3):
        for cell in row:
            cell.value = None

    # Write new data
    current_row = 2
    for player, events in sorted(player_events.items()):
        events_str = ", ".join(sorted(events))
        if player in problematic_players:
            reasons = " ⚠️ " + " | ⚠️ ".join(problematic_players[player])
        else:
            reasons = ""
        sheet.cell(row=current_row, column=1).value = player
        sheet.cell(row=current_row, column=2).value = events_str
        sheet.cell(row=current_row, column=3).value = reasons

        current_row += 1

    book.save(file_path)

    print(f"➡️  Added players to '{sheet_name}'")


from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

def write_event_analysis_to_excel(file_path, event_correlations, round_schedules, event_structures, event_colors, event_players, sheet_name="Event Analysis"):

    try:
        book = load_workbook(file_path)
    except FileNotFoundError:
        print(f"❌ ERROR: The file '{file_path}' was not found. Skipping writing the events analysis to Excel.")
        return
    
    if sheet_name in book.sheetnames:
        del book[sheet_name]
    sheet = book.create_sheet(sheet_name)

    center_aligned = Alignment(horizontal="center", vertical="center")

    def apply_color(cell, event_name, center=True):
        """Applies color formatting based on the provided event_name."""
        if event_name in event_colors:
            hex_bg, text_col = event_colors[event_name]
            clean_hex = hex_bg.replace("#", "")
            cell.fill = PatternFill(start_color=clean_hex, end_color=clean_hex, fill_type="solid")
            cell.font = Font(color="000000" if text_col == "black" else "FFFFFF")
            if center:
                cell.alignment = center_aligned

    all_events = list(event_correlations.keys())

    # --- PART 1: COLLISION MATRIX ---
    # Top Header Row (Remains colored by the column's event for identification)
    for col_idx, col_event in enumerate(all_events, 2):
        cell = sheet.cell(row=1, column=col_idx)
        cell.value = col_event
        apply_color(cell, col_event)

    # Matrix Body
    for row_idx, row_event in enumerate(all_events, 2):
        # Row Header (Column 1)
        row_header_cell = sheet.cell(row=row_idx, column=1)
        row_header_cell.value = row_event
        apply_color(row_header_cell, row_event, center=False)

        # Dictionary of collisions for this specific row
        current_collisions = {other: count for other, count in event_correlations[row_event]}

        for col_idx, col_event in enumerate(all_events, 2):
            target_cell = sheet.cell(row=row_idx, column=col_idx)
            
            # Check if there is a collision (diagonal is always a collision)
            is_self = (row_event == col_event)
            collision_count = current_collisions.get(col_event, 0)

            if is_self or collision_count > 0:
                # Value logic
                if is_self:
                    target_cell.value = len(event_players.get(row_event, []))
                else:
                    target_cell.value = collision_count
                
                # COLOR LOGIC: Always use the Row's color
                apply_color(target_cell, row_event)
            else:
                target_cell.value = None # Keep empty if 0 collisions

    # --- PART 2: EVENT DATA (5 rows below the matrix) ---
    curr_row = len(all_events) + 7 
    
    max_rounds = 0
    if round_schedules:
        max_rounds = max((len(rounds) for rounds in round_schedules.values()), default=0)
    
    headers_base = ["Event", "Player Amount"]
    round_headers = [f"Round {i+1}" for i in range(max_rounds)]
    struct_headers = ["RR", "Elim", "Lower"]
    
    for i, h in enumerate(headers_base + round_headers + struct_headers, 1):
        h_cell = sheet.cell(row=curr_row, column=i)
        h_cell.value = h
        h_cell.font = Font(bold=True)
    
    curr_row += 1

    for event in all_events:
        cell = sheet.cell(row=curr_row, column=1)
        cell.value = event
        apply_color(cell, event, center=False)

        sheet.cell(row=curr_row, column=2).value = len(event_players.get(event, []))
        sheet.cell(row=curr_row, column=2).alignment = center_aligned

        rounds = round_schedules.get(event, [])
        for i, match_count in enumerate(rounds):
            r_cell = sheet.cell(row=curr_row, column=3 + i)
            r_cell.value = match_count
            r_cell.alignment = center_aligned

        if event in event_structures:
            struct = event_structures[event]
            def fmt_list(lst): return ", ".join(map(str, lst)) if lst else ""
            struct_col = 3 + max_rounds
            sheet.cell(row=curr_row, column=struct_col).value = fmt_list(struct.get("round_robin"))
            sheet.cell(row=curr_row, column=struct_col + 1).value = fmt_list(struct.get("elimination"))
            sheet.cell(row=curr_row, column=struct_col + 2).value = fmt_list(struct.get("lower_playoff"))
        curr_row += 1

    book.save(file_path)
    print(f"➡️  Added collision matrix and event data to '{sheet_name}'")

def log_rule_violations_to_excel(file_path, rule_violations, sheet_name="Rule Violations"):
    """
    Logs rule violations to a persistent sheet.
    Keeps all previous data and adds timestamped sections.
    """

    try:
        book = load_workbook(file_path)
    except FileNotFoundError:
        print(f"❌ ERROR: The file '{file_path}' was not found. Skipping writing the rule violations to Excel.")
        return

    # Create sheet if missing
    if sheet_name not in book.sheetnames:
        sheet = book.create_sheet(sheet_name)
        sheet.append(["Date/Time", "Player", "Violation"])
    else:
        sheet = book[sheet_name]

    # Find first empty row
    next_row = sheet.max_row + 2  # leave a blank line
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sheet.cell(row=next_row, column=1).value = f"Log entry at {timestamp}"
    next_row += 1

    # Write violations
    for v in rule_violations:
        player, description = v.split(":", 1) if ":" in v else ("Unknown", v)
        sheet.cell(row=next_row, column=2).value = player.strip()
        sheet.cell(row=next_row, column=3).value = description.strip()
        next_row += 1

    book.save(file_path)
    print(f"➡️  Appended {len(rule_violations)} rule violations to '{sheet_name}'")

'''
def half_hour_slots(start_datetime: str, end_datetime: str):
    """
    Calculates how many 30-minute slots fit between two times (HH:MM format).
    Returns an integer count.
    """
    datetime_format = "%Y-%m-%d %H:%M"
    start = datetime.strptime(start_datetime, datetime_format)
    end = datetime.strptime(end_datetime, datetime_format)
    
    if end <= start:
        raise ValueError("End time must be after start time.")
    
    # Calculate total minutes difference
    slot_length = timedelta(minutes=30)
    
    count = 0
    current = start
    while current <= end:
        count += 1
        current += slot_length
    
    return count
'''
'''
def half_hour_slots(start_time: str, end_time: str):
    """
    Calculates how many 30-minute slots fit between two times (HH:MM format).
    Returns an integer count.
    """
    time_format = "%H:%M"
    
    # These will both default to 1900-01-01
    start = datetime.strptime(start_time, time_format)
    end = datetime.strptime(end_time, time_format)
    
    # Check if end is before start (assuming same day)
    if end <= start:
        print(f"❌ ERROR: End time is before start time! start time: {start_time}, end time: {end_time}. QUITTING!")
        return -1
    
    slot_length = timedelta(minutes=30)
    
    count = 0
    current = start
    # Using <= means it counts the start time AND the end time
    while current <= end:
        count += 1
        current += slot_length
    
    return count
'''

# Calculates the amount of 15 minute times lots that are in a block
def start_time_slots(start_time: str, end_time: str) -> int:
    """
    Calculates how many 15-minute match start times fit between start_time and end_time (inclusive).
    Example: "09:00" to "12:00" -> 13 start slots.
    Example: "12:15" to "16:00" -> 16 start slots.
    """
    time_format = "%H:%M"
    start = datetime.strptime(start_time, time_format)
    end = datetime.strptime(end_time, time_format)
    
    if end < start:
        print(f"❌ ERROR: End time ({end_time}) is before start time ({start_time})! QUITTING!")
        return -1
    
    total_minutes = int((end - start).total_seconds() // 60)
    if total_minutes % 15 != 0:
        print(f"❌ ERROR: Times ({start_time} to {end_time}) must be in 15-minute increments! QUITTING!")
        return -1
        
    return (total_minutes // 15) + 1

# Uses start_time_slots
def parse_day_blocks(day_config, day_name: str):
    """
    Normalizes a day's config, checks that blocks don't overlap,
    and returns (list_of_(number_of_slots, courts), num_of_total_day_slots).
    """
    if not day_config:
        return [], 0

    # Allow both ["09:00", "20:00", 16] and [["09:00", "20:00", 16]]
    if isinstance(day_config[0], str):
        day_config = [day_config]

    blocks = []
    day_total_slots = 0
    time_format = "%H:%M"
    prev_end_dt = None

    for block in day_config:
        start_str, end_str, courts = block[0], block[1], int(block[2])
        start_dt = datetime.strptime(start_str, time_format)
        end_dt = datetime.strptime(end_str, time_format)

        # Ensure blocks on the same day don't overlap or share the same start/end timestamp
        if prev_end_dt is not None and start_dt <= prev_end_dt:
            prev_end_str = prev_end_dt.strftime(time_format)
            print(f"❌ ERROR on {day_name}: Block starting at {start_str} overlaps with previous block ending at {prev_end_str}! (Next block should start at least 15 min later). QUITTING!")
            return None, -1

        slots = start_time_slots(start_str, end_str)
        if slots <= 0:
            return None, -1

        blocks.append((slots, courts))
        day_total_slots += slots
        prev_end_dt = end_dt

    return blocks, day_total_slots

def minizinc_data(print_lists = True):

    file_path = competition_info.file_path
    umpire_chairs = competition_info.umpire_chairs

    event_correlations, event_players, player_events = parse_competition_excel(file_path)
    if not event_correlations and not event_players and not player_events:
        print("❌ ERROR: No data was loaded from the excel (file name likely wrong). QUITTING!")
        return {}
    strong_collision, weak_collision, super_weak_collision = analyze_collisions(event_correlations)
    event_structures, round_schedules = calculate_tournament_structure(event_players, competition_info.pool)
    event_colors = assign_event_colors(event_players.keys(), color_bank)

    '''
    # Time slots that can hold matches (15 min)
    #day_zero = half_hour_slots("2025-10-10 09:00", "2025-10-10 21:00") * 2 - 1 
    day_zero = 0#half_hour_slots("2025-10-10 18:00", "2025-10-10 21:00") * 2 - 1
    if competition_info.friday:
        day_zero = half_hour_slots(competition_info.friday_start_time, competition_info.friday_end_time) * 2 - 1

    day_one = 0
    if competition_info.saturday:
        day_one = half_hour_slots(competition_info.saturday_start_time, competition_info.saturday_end_time) * 2 - 1
    
    day_two_1 = 0#half_hour_slots("2025-10-10 09:00", "2025-10-10 10:00") * 2 - 1
    if competition_info.sunday_1:
        day_two_1 = half_hour_slots(competition_info.sunday_start_time_1, competition_info.sunday_end_time_1) * 2 - 1

    day_two_2 = 0#half_hour_slots("2025-10-10 10:30", "2025-10-10 14:00") * 2 - 1
    if competition_info.sunday_2:
        day_two_2 = half_hour_slots(competition_info.sunday_start_time_2, competition_info.sunday_end_time_2) * 2 - 1

    day_two_3 = 0#half_hour_slots("2025-10-10 14:30", "2025-10-10 19:00") * 2 - 1
    if competition_info.sunday_3:
        day_two_3 = half_hour_slots(competition_info.sunday_start_time_3, competition_info.sunday_end_time_3) * 2 - 1

    if day_zero < 0 or day_one < 0 or day_two_1 < 0 or day_two_2 < 0 or day_two_3 < 0:
        print("❌ ERROR: the timeslots could not be counted! QUITTING!")
        return


    # The tree different time slots based on how many fields are available
    timeslots_12_fields = day_zero + day_one + day_two_1
    timeslots_8_fields = day_two_2
    timeslots_6_fields = day_two_3

    # amount of fields
    fields_1 = competition_info.friday_saturday_sunday_1_courts
    fields_2 = competition_info.sunday_courts_2
    fields_3 = competition_info.sunday_courts_3

    # amount of matchslots
    matchslots_1 = timeslots_12_fields * fields_1
    matchslots_2 = timeslots_8_fields * fields_2
    matchslots_3 = timeslots_6_fields * fields_3

    # the last slots of day 0 and last of day 1
    last_of_day0 = day_zero
    last_of_day1 = (day_zero + day_one)

    '''
    fri_blocks, num_of_slots_day_zero = parse_day_blocks(competition_info.friday, "Friday")
    sat_blocks, num_of_slots_day_one  = parse_day_blocks(competition_info.saturday, "Saturday")
    sun_blocks, num_of_slots_day_two  = parse_day_blocks(competition_info.sunday, "Sunday")

    if num_of_slots_day_zero < 0 or num_of_slots_day_one < 0 or num_of_slots_day_two < 0:
        print("❌ ERROR: The timeslots could not be counted! QUITTING!")
        return {}
    
    all_blocks = fri_blocks + sat_blocks + sun_blocks
    if not all_blocks:
        print("❌ ERROR: No active competition days configured! QUITTING!")
        return {}

    # Dynamic block lists for MiniZinc
    nof_blocks = len(all_blocks)
    block_timeslots = [b[0] for b in all_blocks]
    block_fields = [b[1] for b in all_blocks]
    block_matchslots = [b[0] * b[1] for b in all_blocks]

    # 1-indexed start slot for each block in MiniZinc
    block_starts = []
    current_slot = 1
    for slots in block_timeslots:
        block_starts.append(current_slot)
        current_slot += slots

    # Last slots of Day 0 (Friday) and Day 1 (Saturday)
    last_of_day0 = num_of_slots_day_zero
    last_of_day1 = num_of_slots_day_zero + num_of_slots_day_one

    # 2d list where index is class and element is list of matches (size = rounds)
    elimination_data_minizinc = [i[1] for i in round_schedules.items()]

    event_number_dict = {i[1]: i[0] + 1 for i in enumerate(round_schedules.keys())}

    number_event_dict = {x[1]: x[0] for x in event_number_dict.items()}

    event_to_number = lambda x: event_number_dict[x]

    # collisions paired up
    seen_pairs = set()
    strong_collision_minizinc = []

    for i, (_, values) in enumerate(strong_collision.items()):
        for v in values:
            a = i + 1
            b = event_to_number(v)
            key = tuple(sorted((a, b)))  # Ensure (1,2) and (2,1) are treated the same
            if key not in seen_pairs:
                seen_pairs.add(key)
                strong_collision_minizinc.extend([a, b])
    
    seen_pairs = set()
    weak_collision_minizinc = []
    for i, (_, values) in enumerate(weak_collision.items()):
        for v in values:
            a = i + 1
            b = event_to_number(v[0])   # v is a tuple of event name and amount of colliding players
            key = tuple(sorted((a, b)))
            if key not in seen_pairs:
                seen_pairs.add(key)
                weak_collision_minizinc.extend([a, b])
    
    seen_pairs = set()
    super_weak_collision_minizinc = []
    for i, (_, values) in enumerate(super_weak_collision.items()):
        for v in values:
            a = i + 1
            b = event_to_number(v[0])   # v is a tuple of event name and amount of colliding players
            key = tuple(sorted((a, b)))
            if key not in seen_pairs:
                seen_pairs.add(key)
                super_weak_collision_minizinc.extend([a, b])


    matches_minizic = []
    V_matches = []
    single_V_matches = []
    A_matches = []
    junior_matches = []
    matches_rest = []
    is_elite_match = False
    is_single_V_match = False
    is_A_match = False
    is_junior_match = False
    class_index = []
    round_index = []
    first_singles_round_M = []
    first_doubles_mixt_round_M = []
    first_singles_round_W = []
    first_doubles_mixt_round_W = []

    # fix so it counts!!!
    # start_of_V_classes = 0 good to have later
    nof_V_classes = 0
    nof_single_V_classes = 0
    nof_A_classes = 0
    nof_junior_classes = 0
    nof_rest_classes = 0

    
    # Iterate through the event names in the event_players dictionary
    for event_name in event_players.keys():
        # 1. Check for V classes (ending with V)
        if event_name.endswith("V"):
            nof_V_classes += 1
            if "S" in event_name:
                nof_single_V_classes += 1
            
        # 2. Check for A classes (ending with A)
        elif event_name.endswith("A"):
            nof_A_classes += 1
            
        # 3. Check for Junior classes (containing U9, U11, or U13)
        # Using 'in' is safer here in case the name is 'BS U13' or 'BS U13 '
        elif any(jr in event_name for jr in ["U9", "U11", "U13"]):
            nof_junior_classes += 1
            
        # 4. Everything else
        else:
            nof_rest_classes += 1
    

    # Initialize index trackers for when the special classes start and end
    V_range = {"start": 0, "end": 0}
    single_V_range = {"start": 0, "end": 0}
    A_range = {"start": 0, "end": 0}
    junior_range = {"start": 0, "end": 0}

    # Initialize index trackers for when the special rounds start and end
    V_rounds_range = {"start": 0, "end": 0}
    single_V_rounds_range = {"start": 0, "end": 0}
    A_rounds_range = {"start": 0, "end": 0}
    junior_rounds_range = {"start": 0, "end": 0}

    count = 1
    class_nr = 1
    round_counter = 1
    

    for outer in elimination_data_minizinc:
        is_elite_match = False
        is_single_V_match = False
        is_A_match = False
        is_junior_match = False

        class_index.append(count)

        current_event_name = number_event_dict[class_nr]

        if current_event_name.endswith("V"):
            is_elite_match = True
            if V_range["start"] == 0: V_range["start"] = class_nr
            V_range["end"] = class_nr
            if "S" in current_event_name:
                is_single_V_match = True
                if single_V_range["start"] == 0: single_V_range["start"] = class_nr
                single_V_range["end"] = class_nr

        elif current_event_name.endswith("A"):
            is_A_match = True
            if A_range["start"] == 0: A_range["start"] = class_nr
            A_range["end"] = class_nr

        elif any(current_event_name.endswith(s) for s in ["9", "11", "13"]):
            is_junior_match = True
            if junior_range["start"] == 0: junior_range["start"] = class_nr
            junior_range["end"] = class_nr
        
        for inner in outer:
            # --- Round Tracking Logic ---
            if is_elite_match:
                if V_rounds_range["start"] == 0: V_rounds_range["start"] = round_counter
                V_rounds_range["end"] = round_counter
                if is_single_V_match:
                    if single_V_rounds_range["start"] == 0: single_V_rounds_range["start"] = round_counter
                    single_V_rounds_range["end"] = round_counter
            elif is_A_match:
                if A_rounds_range["start"] == 0: A_rounds_range["start"] = round_counter
                A_rounds_range["end"] = round_counter
            elif is_junior_match:
                if junior_rounds_range["start"] == 0: junior_rounds_range["start"] = round_counter
                junior_rounds_range["end"] = round_counter

            round_counter += 1

            round_index.append(count)
            for _ in range(0, inner):
                if is_elite_match:
                    V_matches.append(count)
                    if is_single_V_match:
                        single_V_matches.append(count)
                elif is_A_match:
                    A_matches.append(count)
                elif is_junior_match:
                    junior_matches.append(count)
                else:
                    matches_rest.append(count)
                matches_minizic.append(count)
                count += 1
        class_nr += 1

    # --- Validation Logic ---
    # Helper to validate: (end - start + 1) should equal the count we found earlier
    validations = [
        ("V", V_range, nof_V_classes),
        ("Single V", single_V_range, nof_single_V_classes),
        ("A", A_range, nof_A_classes),
        ("Junior", junior_range, nof_junior_classes)
    ]

    for label, r, expected in validations:
        actual = (r["end"] - r["start"] + 1) if r["start"] != 0 else 0
        if actual != expected:
            print(f"‼️ WARNING: {label} class mismatch! Range {r['start']}-{r['end']} "
                f"suggests {actual} classes, but {expected} were expected.")
    # --- Validation Logic ---
    
    # number of rounds for the special classes
    nof_junior_rounds = sum(len(arr) for arr in elimination_data_minizinc[junior_range["start"]-1 : junior_range["end"]])
    nof_V_rounds = sum(len(arr) for arr in elimination_data_minizinc[V_range["start"]-1 : V_range["end"]])
    nof_single_V_rounds = sum(len(arr) for arr in elimination_data_minizinc[single_V_range["start"]-1 : single_V_range["end"]])
    nof_A_rounds = sum(len(arr) for arr in elimination_data_minizinc[A_range["start"]-1 : A_range["end"]])
    
    if print_lists:
        print("NUMBER OF V EVENTS:")
        print(nof_V_classes, end='\n')
        print("NUMBER OF SINGLE V EVENTS:")
        print(nof_single_V_classes, end='\n')
        print("NUMBER OF A EVENTS:")
        print(nof_A_classes, end='\n')
        print("NUMBER OF JUNIOR (U9, U11, U13) EVENTS:")
        print(nof_junior_classes, end='\n')

        print("NUMBER OF REMAINING EVENTS:")
        print(nof_rest_classes, end='\n\n')

        print("NUMBER OF V ROUNDS:")
        print(nof_V_rounds, end='\n')
        print("NUMBER OF SINGLE V ROUNDS:")
        print(nof_single_V_rounds, end='\n')
        print("NUMBER OF A ROUNDS:")
        print(nof_A_rounds, end='\n')
        print("NUMBER OF JUNIOR (U9, U11, U13) ROUNDS:")
        print(nof_junior_rounds, end='\n\n')

        print("\n\n")


        print("CLASS NUMBERS:")
        print(number_event_dict, end="\n\n")

        print("ELIMINATION DATA:")
        print(elimination_data_minizinc, end='\n\n')
        print("CLASS INDEXES:")
        print(class_index, end='\n\n')
        print("ROUND INDEXES:")
        print(round_index, end='\n\n')
        print("MATCHES:")
        print(matches_minizic, end='\n\n')
        #print("UMPIRE MATCHES:")
        #print(umpire_matches, end='\n\n')
        #print("MATCHES WITHOUT UMPIRE MATCHES:")
        #print(matches_no_umpire, end='\n\n')
        print("STRONG COLLISION:")
        print(strong_collision, end='\n\n')
        print("STRONG COLLISION MINIZINC:")
        print(strong_collision_minizinc, end="\n\n")
        print("WEAK COLLISION:")
        print(weak_collision, end="\n\n")
        print("WEAK COLLISION MINIZINC:")
        print(weak_collision_minizinc, end="\n\n")
        print("SUPER WEAK COLLISION:")
        print(super_weak_collision, end="\n\n")
        print("SUPER WEAK COLLISION MINIZINC:")
        print(super_weak_collision_minizinc, end="\n\n")
        #print("AMOUNT OF ELITE ROUNDS")
        #print(nof_elit_rounds, end="\n\n")
        print("BLOCKS (1 block is timeslots and courts)")
        print(all_blocks, end="\n\n")
        print("BLOCK MATCHSLOTS (timeslots * courts)")
        print(block_matchslots, end="\n\n")

    data = {
        "file_path": file_path,

        "umpire_chairs": umpire_chairs,

        "strong_collision": strong_collision_minizinc,
        "weak_collision": weak_collision_minizinc,
        "super_weak_collision": super_weak_collision_minizinc,

        "class_index": class_index,
        "round_index": round_index,

        "nof_blocks": nof_blocks,
        "block_starts": block_starts,
        "block_timeslots": block_timeslots,
        "block_fields": block_fields,
        "block_matchslots": block_matchslots,

        "last_of_day1": last_of_day1,
        "last_of_day0": last_of_day0,

        "nof_V_classes": nof_V_classes,
        "nof_single_V_classes": nof_single_V_classes,
        "nof_A_classes": nof_A_classes,
        "nof_junior_classes": nof_junior_classes,

        "V_start": V_range["start"],
        "V_end": V_range["end"],
        "single_V_start": single_V_range["start"],
        "single_V_end": single_V_range["end"],
        "A_start": A_range["start"],
        "A_end": A_range["end"],
        "junior_start": junior_range["start"],
        "junior_end": junior_range["end"],

        "nof_V_rounds": nof_V_rounds,
        "nof_single_V_rounds": nof_single_V_rounds,
        "nof_A_rounds": nof_A_rounds,
        "nof_junior_rounds": nof_junior_rounds,

        "V_rounds_start": V_rounds_range["start"],
        "V_rounds_end": V_rounds_range["end"],
        "single_V_rounds_start": single_V_rounds_range["start"],
        "single_V_rounds_end": single_V_rounds_range["end"],
        "A_rounds_start": A_rounds_range["start"],
        "A_rounds_end": A_rounds_range["end"],
        "junior_rounds_start": junior_rounds_range["start"],
        "junior_rounds_end": junior_rounds_range["end"],
        
        "matches": matches_minizic,
        "V_matches": V_matches,
        "single_V_matches": single_V_matches,
        "A_matches": A_matches,
        "junior_matches": junior_matches,
        "matches_rest": matches_rest,

        "event_colors": event_colors,
    }

    return data

def main():
    print("\n\n----Running parse_excel.py----\n")
    
    file_path = competition_info.file_path
    #file_path = "Entries Dubbelturnering 2026 1.1. .xlsx"#"Junnu entries.xlsx"  # <-- your Excel file
    print(f"Reading from and writing to '{file_path}'\n")
    print("Performed action:")
    event_correlations, event_players, player_events = parse_competition_excel(file_path)
    if not event_correlations and not event_players and not player_events:
        print("❌ ERROR: No data was loaded from the excel (file name likely wrong). QUITTING!")
        print("\n----Done due to error----\n")
        return
    event_structures, round_schedules = calculate_tournament_structure(event_players, competition_info.pool)
    rule_violations = check_rule_violations(player_events)
    problematic_players = find_problematic_players(player_events, event_players, event_correlations)
    event_colors = assign_event_colors(event_players.keys(), color_bank)
    write_player_events_to_excel(file_path, player_events, problematic_players, sheet_name="Player Events")
    write_event_analysis_to_excel(file_path, event_correlations, round_schedules, event_structures, event_colors, event_players, sheet_name="Event Analysis")
    log_rule_violations_to_excel(file_path, rule_violations, sheet_name="Rule Violations")
    

    # Print all players for all events
    
    print("\n=== Event -> Players ===")
    for event, players in event_players.items():
        print(f"{event}: {sorted(players)}")
    
    # Print the amount of entries for all events
    '''
    print("\n=== Event -> num Players ===")
    for event, participants in participants_structure.items():
        print(f"{event}: {participants}")
    '''
    # Print all events for all player
    
    print("\n=== Player -> Events ===")
    for player, events in player_events.items():
        print(f"{player}: {sorted(events)}")
    
    # Print all collitions for an event (with wight)
    
    print("\n=== Event Correlations ===")
    for event, correlations in event_correlations.items():
        print(f"\nEvent: {event}")
        for other_event, count in correlations:
            print(f"  {other_event}: {count}")
    
    # Print strong collisions for an event
    '''
    print("\n=== Events colliding with same last character ===")
    for event, others in strong_collision.items():
        print(f"{event}: {others}")
    
    # Print weak collitions for an event (with wight)
    
    print("\n=== Events colliding with different last character (with counts) ===")
    for event, others in weak_collision.items():
        print(f"{event}: {others}")
    '''
    # Print mathces per elimination round per event
    
    print("\n=== Event structures ===")
    for event, rounds in event_structures.items():
        print(f"{event}: {rounds}")
    
    print("\n=== Rounds and matches per round ===")
    for event, rounds in round_schedules.items():
        print(f"{event}: {rounds}")
    
    
    # Print warnings related to players entries
    
    for warning in rule_violations:
        print(warning)
    
    # Print summary
    
    print("\n=== Problematic Players ===")
    if not problematic_players:
        print("✅ No problematic players found.")
    else:
        for player, reasons in problematic_players.items():
            print(f"⚠️ {player}:")
            for reason in reasons:
                print(f"   - {reason}")
    
    print("\n----Done----\n")

if __name__ == "__main__":
    #minizinc_data()
    main()